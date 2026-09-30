"""Pembaca .docx alat uji — bug 10 (docs/perbaiki bug.md).

Pembaca lama membaca sel gabungan sekali untuk tiap kolom yang ditempatinya
dan melewatkan tabel bersarang. PMK 119 terbaca 127 pasal padahal 64, dan
F2-004 melaporkan 63 nomor pasal ganda yang tidak pernah ada di Word.
"""

from pathlib import Path

import pytest
from docx import Document

from tools import baca_docx as B


def _simpan(doc, tmp_path: Path) -> Path:
    path = tmp_path / "uji.docx"
    doc.save(str(path))
    return path


def test_sel_gabungan_mendatar_dibaca_sekali(tmp_path):
    doc = Document()
    t = doc.add_table(rows=1, cols=3)
    gabung = t.cell(0, 0).merge(t.cell(0, 2))
    gabung.text = "Pasal 5"
    paragraf, _ = B.baca_docx(_simpan(doc, tmp_path))
    assert [p["teks"] for p in paragraf].count("Pasal 5") == 1


def test_sel_gabungan_menurun_tidak_menggandakan_teks(tmp_path):
    doc = Document()
    t = doc.add_table(rows=3, cols=2)
    t.cell(0, 0).merge(t.cell(2, 0)).text = "NO."
    t.cell(0, 1).text = "satu"
    t.cell(1, 1).text = "dua"
    paragraf, _ = B.baca_docx(_simpan(doc, tmp_path))
    assert [p["teks"] for p in paragraf].count("NO.") == 1


def test_tabel_bersarang_ikut_terbaca_dan_letaknya_tercatat(tmp_path):
    doc = Document()
    luar = doc.add_table(rows=1, cols=2)
    luar.cell(0, 0).text = "luar"
    dalam = luar.cell(0, 1).add_table(rows=1, cols=2)
    dalam.cell(0, 0).text = "Uraian"
    dalam.cell(0, 1).text = "Jumlah"
    paragraf, _ = B.baca_docx(_simpan(doc, tmp_path))
    uraian = next(p for p in paragraf if p["teks"] == "Uraian")
    jumlah = next(p for p in paragraf if p["teks"] == "Jumlah")
    assert uraian["tabel"] == jumlah["tabel"] != -1
    assert (uraian["baris"], uraian["sel"], jumlah["sel"]) == (0, 0, 1)
    assert next(p for p in paragraf if p["teks"] == "luar")["tabel"] != uraian["tabel"]


def test_paragraf_di_luar_tabel_tanpa_letak_tabel(tmp_path):
    doc = Document()
    doc.add_paragraph("Pasal 1")
    paragraf, _ = B.baca_docx(_simpan(doc, tmp_path))
    p = next(p for p in paragraf if p["teks"] == "Pasal 1")
    assert (p["tabel"], p["baris"], p["sel"], p["letak_pasti"]) == (-1, -1, -1, True)


def test_tabel_raksasa_jadi_kerangka_dan_letak_sesudahnya_tidak_pasti(tmp_path, monkeypatch):
    monkeypatch.setattr(B, "BATAS_BARIS_RAKSASA", 3)
    doc = Document()
    doc.add_paragraph("sebelum")
    t = doc.add_table(rows=6, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Desa", "Jumlah"
    for i in range(1, 6):
        t.cell(i, 0).text, t.cell(i, 1).text = f"desa {i}", str(i * 10)
    doc.add_paragraph("sesudah")
    paragraf, kerangka = B.baca_docx(_simpan(doc, tmp_path))
    teks = [p["teks"] for p in paragraf]
    assert "desa 3" not in teks  # isi tabel raksasa tidak dibaca per paragraf
    assert len(kerangka) == 1
    k = kerangka[0]
    assert (k["jumlah_baris"], k["jumlah_kolom"]) == (6, 2)
    assert k["contoh"][0] == ["Desa", "Jumlah"]
    assert paragraf[k["sesudah_paragraf"]]["teks"] == "sebelum"
    sesudah = next(p for p in paragraf if p["teks"] == "sesudah")
    assert sesudah["letak_pasti"] is False
    assert next(p for p in paragraf if p["teks"] == "sebelum")["letak_pasti"] is True


def test_index_berurutan(tmp_path):
    doc = Document()
    for i in range(3):
        doc.add_paragraph(f"p{i}")
    paragraf, _ = B.baca_docx(_simpan(doc, tmp_path))
    assert [p["index"] for p in paragraf] == list(range(len(paragraf)))


@pytest.mark.parametrize("teks", ["a\tb", "baris\nbaru"])
def test_teks_sama_dengan_python_docx(tmp_path, teks):
    doc = Document()
    doc.add_paragraph(teks)
    path = _simpan(doc, tmp_path)
    paragraf, _ = B.baca_docx(path)
    asli = [p.text for p in Document(str(path)).paragraphs]
    assert [p["teks"] for p in paragraf] == asli
